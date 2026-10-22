"""Sprint 4: US17 (export), US18 (retraining), US19 (smart replies), US20 (FAQ suggestions), US21 (knowledge base).

Each USxQATest Given/When/Then scenario has at least one test here (see docs/SPRINT4_PLAN.md for the map).
"""
import io
import random
import time
from datetime import timedelta

import joblib
import pytest
from openpyxl import load_workbook

from app.db import utcnow
from app.ml import classifier, dataset, embeddings, pii, training
from app.models import (AIFeedback, AuditLog, FAQArticle, ModelRegistry, Notification, ReplyTemplate, Request,
                        SuggestionEvent, Unit, User)
from app.seed import seed_knowledge_base
from conftest import login, set_online

API = "/api/v1"


def _submit(client, text, email="student@qoldau.edu.kz", subject="Request", **extra):
    h = login(client, email)
    r = client.post(f"{API}/requests", headers=h, json={"subject": subject, "description": text, **extra})
    assert r.status_code == 201, r.text
    return r.json(), h


def _resolve(client, req_id, agent_email, reply):
    h = login(client, agent_email)
    assert client.post(f"{API}/requests/{req_id}/messages", headers=h, json={"body": reply}).status_code == 201
    r = client.patch(f"{API}/requests/{req_id}/status", headers=h, json={"status": "Resolved"})
    assert r.status_code == 200, r.text
    return h


@pytest.fixture()
def kb(db):
    """Seeded knowledge base (20 topics x KZ/RU/EN, 15 templates), indexed."""
    seed_knowledge_base(db, random.Random(4))
    embeddings.reindex_all(db)
    db.commit()
    return db


@pytest.fixture()
def restore_model():
    yield
    classifier.reset()  # later tests use the default model again


# ======================= US21 Knowledge Base =======================
def test_us21_unit_head_publishes_faq_in_three_languages(client, db):
    """Given the Registrar unit head, When they create 'Ordering a transcript' in KZ/RU/EN and publish,
    Then status = Published, version 1 is saved, And within 1 min student suggestions return it."""
    h = login(client, "head.registrar@qoldau.edu.kz")
    texts = {
        "en": ("Ordering a transcript", "<p>Order an official transcript in Qoldau; it is ready in 2 working days.</p>"),
        "ru": ("Как заказать транскрипт", "<p>Закажите официальный транскрипт в Qoldau, он готов через 2 рабочих дня.</p>"),
        "kz": ("Транскриптке тапсырыс беру", "<p>Ресми транскриптке Qoldau арқылы тапсырыс беріңіз, 2 жұмыс күнінде дайын.</p>"),
    }
    ids = {}
    for lang, (title, body) in texts.items():
        r = client.post(f"{API}/admin/faq", headers=h, json={"slug": "ordering-transcript", "title": title, "body": body,
                                                             "language": lang, "category_id": 6, "tags": ["transcript"]})
        assert r.status_code == 201, r.text
        assert r.json()["status"] == "draft" and r.json()["unit_name"] == "Registrar Office"
        ids[lang] = r.json()["id"]
        p = client.post(f"{API}/admin/faq/{ids[lang]}/publish", headers=h)
        assert p.status_code == 200 and p.json()["status"] == "published"
        versions = client.get(f"{API}/admin/faq/{ids[lang]}/versions", headers=h).json()
        assert [v["version"] for v in versions] == [1]
    assert db.query(AuditLog).filter_by(entity="faq", action="publish").count() == 3

    hs = login(client, "student@qoldau.edu.kz")
    for lang, query in (("en", "how do I get a transcript"), ("ru", "как получить транскрипт"),
                        ("kz", "транскрипт қалай алуға болады")):
        out = client.post(f"{API}/ai/similar", headers=hs, json={"text": query}).json()
        assert f"faq:{ids[lang]}" in [i["id"] for i in out["faq"]], (lang, out)


def test_us21_agent_proposes_faq_and_cannot_publish(client, db):
    """Given an agent resolved a request, When they click 'Propose FAQ', Then a Draft is created and the
    unit head is notified, And the agent cannot publish it (HTTP 403)."""
    set_online(db, "agent1.it@qoldau.edu.kz")
    req, _ = _submit(client, "I cannot log in to Moodle, the password reset link does not arrive")
    _resolve(client, req["id"], "agent1.it@qoldau.edu.kz",
             "Hello Aruzhan, I reset your Moodle password; use the link sent to aruzhan@qoldau.edu.kz.")
    ha = login(client, "agent1.it@qoldau.edu.kz")
    r = client.post(f"{API}/requests/{req['id']}/propose-faq", headers=ha, json={})
    assert r.status_code == 201, r.text
    art = r.json()
    assert art["status"] == "draft" and art["source_request_id"] == req["id"]
    assert "Aruzhan" not in art["body"] and "aruzhan@" not in art["body"]  # proposal is PII-masked
    head = db.query(User).filter_by(email="head.it@qoldau.edu.kz").one()
    assert db.query(Notification).filter_by(user_id=head.id, type="faq_proposed").count() == 1

    assert client.post(f"{API}/admin/faq/{art['id']}/publish", headers=ha).status_code == 403
    other_head = login(client, "head.library@qoldau.edu.kz")
    assert client.post(f"{API}/admin/faq/{art['id']}/publish", headers=other_head).status_code == 403
    # Drafts never reach students.
    hs = login(client, "student2@qoldau.edu.kz")
    out = client.post(f"{API}/ai/similar", headers=hs, json={"text": "cannot log in to Moodle password reset"}).json()
    assert f"faq:{art['id']}" not in [i["id"] for i in out["faq"]]
    assert client.post(f"{API}/admin/faq/{art['id']}/publish", headers=login(client, "head.it@qoldau.edu.kz")).status_code == 200


def test_us21_rich_text_is_sanitised_and_every_edit_is_versioned(client, db):
    h = login(client, "admin@qoldau.edu.kz")
    body = '<p onclick="steal()">Hi<script>alert(1)</script> <a href="javascript:alert(1)">x</a></p><img src=x onerror=1>'
    a = client.post(f"{API}/admin/faq", headers=h, json={"title": "Test article", "body": body, "language": "en"}).json()
    assert "<script" not in a["body"] and "onclick" not in a["body"] and "javascript:" not in a["body"]
    assert "<img" not in a["body"]
    upd = client.put(f"{API}/admin/faq/{a['id']}", headers=h,
                     json={"title": "Test article", "body": "<p>Updated <strong>text</strong></p>", "language": "en"})
    assert upd.json()["version"] == 2
    assert [v["version"] for v in client.get(f"{API}/admin/faq/{a['id']}/versions", headers=h).json()] == [2, 1]
    assert client.get(f"{API}/admin/faq", headers=login(client, "student@qoldau.edu.kz")).status_code == 403


def test_us21_save_reply_as_template_uses_placeholders(client, db):
    set_online(db, "agent1.it@qoldau.edu.kz")
    req, _ = _submit(client, "My university email account is locked after too many attempts")
    student = db.query(User).filter_by(email="student@qoldau.edu.kz").one()
    _resolve(client, req["id"], "agent1.it@qoldau.edu.kz", f"Hello {student.name}, your account {req['id']} is unlocked.")
    t = client.post(f"{API}/requests/{req['id']}/save-template", headers=login(client, "agent1.it@qoldau.edu.kz"),
                    json={}).json()
    assert "{student_name}" in t["body"] and "{request_id}" in t["body"] and student.name not in t["body"]


# ======================= US20 FAQ & similar-request suggestions =======================
def test_us20_faq_suggestions_while_typing(client, kb):
    """Given the New Request form, When the student types 'Как получить справку с места учебы',
    Then within 800 ms the 'Enrollment certificate' FAQ (Registrar Office) is in the top-3."""
    hs = login(client, "student@qoldau.edu.kz")
    t = time.perf_counter()
    out = client.post(f"{API}/ai/similar", headers=hs, json={"text": "Как получить справку с места учебы"}).json()
    assert time.perf_counter() - t < 0.8
    slugs = [kb.get(FAQArticle, int(i["id"].split(":")[1])).slug for i in out["faq"]]
    assert "enrollment-certificate" in slugs[:3]
    top = out["faq"][slugs.index("enrollment-certificate")]
    assert top["unit"] == "Registrar Office" and top["language"] == "ru" and "<p>" in top["body"]  # opens inline


def test_us20_self_service_deflection(client, kb):
    """Given an FAQ answer is shown, When the student clicks 'This solved my problem',
    Then no request is created and a deflection event is recorded, And the deflection rate increases."""
    ha = login(client, "admin@qoldau.edu.kz")
    before = client.get(f"{API}/analytics/dashboard", headers=ha).json()["kpis"]
    hs = login(client, "student@qoldau.edu.kz")
    out = client.post(f"{API}/ai/similar", headers=hs, json={"text": "how to get a transcript"}).json()
    item = out["faq"][0]["id"]
    for action in ("shown", "opened", "deflected"):
        r = client.post(f"{API}/suggestions/events", headers=hs,
                        json={"session_id": "s-1", "type": "faq", "item_ids": [item], "action": action})
        assert r.status_code == 200
    assert kb.query(Request).count() == 0
    assert kb.query(SuggestionEvent).filter_by(action="deflected").count() == 1
    after = client.get(f"{API}/analytics/dashboard", headers=ha).json()["kpis"]
    assert after["deflections"] == before["deflections"] + 1
    assert (after["deflection_rate"] or 0) > (before["deflection_rate"] or 0)
    # Opening counted as a view; agents cannot fake deflections.
    assert kb.get(FAQArticle, int(item.split(":")[1])).views >= 1
    r = client.post(f"{API}/suggestions/events", headers=login(client, "agent1.it@qoldau.edu.kz"),
                    json={"type": "faq", "item_ids": [item], "action": "deflected"})
    assert r.status_code == 403


def test_us20_submitting_after_suggestions_saves_them_with_the_request(client, kb):
    hs = login(client, "student@qoldau.edu.kz")
    out = client.post(f"{API}/ai/similar", headers=hs, json={"text": "my scholarship did not arrive"}).json()
    ids = [i["id"] for i in out["faq"]]
    client.post(f"{API}/suggestions/events", headers=hs,
                json={"session_id": "s-2", "type": "faq", "item_ids": ids, "action": "shown"})
    req, _ = _submit(client, "My scholarship did not arrive this month, please check", suggestion_session="s-2")
    saved = kb.get(Request, req["id"])
    assert saved.suggestions_shown == sorted(ids)
    assert kb.query(SuggestionEvent).filter_by(session_id="s-2", request_id=req["id"]).count() == len(ids)


def test_us20_never_draft_or_other_students_data(client, db):
    set_online(db, "agent1.registrar@qoldau.edu.kz", "agent2.registrar@qoldau.edu.kz", "agent3.registrar@qoldau.edu.kz")
    other = db.query(User).filter_by(email="student3@qoldau.edu.kz").one()
    req, _ = _submit(client, f"I need an enrollment certificate for the bank, my phone is +7 701 555 12 34, "
                             f"email {other.email}", email=other.email, subject="Enrollment certificate for the bank")
    agent = db.get(Request, req["id"]).assignee.email
    _resolve(client, req["id"], agent, f"Hello {other.name.split()[0]}, your certificate is ready at the Registrar Office.")
    hi = login(client, "admin@qoldau.edu.kz")
    client.post(f"{API}/admin/faq", headers=hi, json={"title": "Enrollment certificate draft",
                                                      "body": "<p>Draft about the enrollment certificate for banks.</p>"})
    out = client.post(f"{API}/ai/similar", headers=login(client, "student@qoldau.edu.kz"),
                      json={"text": "enrollment certificate for the bank"}).json()
    assert out["faq"] == []  # only a draft exists
    assert out["cases"], out
    blob = str(out)
    for secret in (other.name, other.name.split()[0], other.email, "555 12 34", req["id"]):
        assert secret not in blob


# ======================= US19 Smart replies =======================
def test_us19_agent_uses_a_smart_reply(client, kb):
    """Given an IT agent opens a Russian Moodle request, When the Smart Reply panel opens, Then 3 Russian
    suggestions with source and score appear within 1.5 s, And Insert fills placeholders without sending."""
    set_online(kb, "agent1.it@qoldau.edu.kz", "agent2.it@qoldau.edu.kz", "agent3.it@qoldau.edu.kz")
    req, _ = _submit(client, "Не могу войти в Moodle, пишет неверный пароль", subject="Не могу войти в Moodle")
    r = kb.get(Request, req["id"])
    agent = r.assignee
    ha = login(client, agent.email)
    t = time.perf_counter()
    out = client.get(f"{API}/ai/suggestions/{req['id']}", headers=ha).json()
    assert time.perf_counter() - t < 1.5
    assert out["language"] == "ru" and len(out["suggestions"]) == 3
    assert all(s["language"] == "ru" and s["source"] in ("template", "similar_case", "kb") for s in out["suggestions"])
    first = out["suggestions"][0]
    assert first["source"] == "template" and "Moodle" in first["title"]
    assert r.student.name in first["text"] and agent.name in first["text"] and "{" not in first["text"]
    assert kb.query(SuggestionEvent).filter_by(request_id=r.id, type="smart_reply", action="shown").count() == 3
    # Inserting never sends: the thread is still empty until the agent clicks Send.
    assert client.get(f"{API}/requests/{req['id']}/messages", headers=ha).json() == []
    tid = int(first["source_id"].split(":")[1])
    before = kb.get(ReplyTemplate, tid).usage_count
    client.post(f"{API}/suggestions/events", headers=ha, json={"request_id": req["id"], "type": "smart_reply",
                                                               "item_ids": [first["source_id"]], "action": "edited"})
    kb.expire_all()
    assert kb.get(ReplyTemplate, tid).usage_count == before + 1
    assert client.get(f"{API}/ai/suggestions/{req['id']}", headers=login(client, "student@qoldau.edu.kz")).status_code == 403


def test_us19_suggestions_mask_pii_and_exclude_internal_notes(client, db):
    """Given a similar resolved request contains another student's name and internal notes, When suggestions
    are generated, Then personal data are masked and notes excluded, And used/edited/ignored is recorded."""
    set_online(db, "agent1.it@qoldau.edu.kz")
    other = db.query(User).filter_by(email="student5@qoldau.edu.kz").one()
    old, _ = _submit(client, "I cannot connect to the campus Wi-Fi on my laptop", email=other.email)
    ha = login(client, "agent1.it@qoldau.edu.kz")
    client.post(f"{API}/requests/{old['id']}/messages", headers=ha,
                json={"body": "SECRET-NOTE student has unpaid fees", "is_internal": True})
    _resolve(client, old["id"], "agent1.it@qoldau.edu.kz",
             f"Hello {other.name}, connect to SDU-Students with PEAP and your e-mail {other.email}.")
    new, _ = _submit(client, "Wi-Fi does not connect on my laptop in the dorm")
    out = client.get(f"{API}/ai/suggestions/{new['id']}", headers=ha).json()
    cases = [s for s in out["suggestions"] if s["source"] == "similar_case"]
    assert cases, out
    blob = str(out)
    assert "SECRET-NOTE" not in blob and other.email not in blob and other.name not in blob
    assert db.get(Request, new["id"]).student.name in cases[0]["text"]  # the current student is greeted instead
    for action in ("used", "ignored"):
        client.post(f"{API}/suggestions/events", headers=ha, json={"request_id": new["id"], "type": "smart_reply",
                                                                   "item_ids": [cases[0]["source_id"]], "action": action})
    assert {e.action for e in db.query(SuggestionEvent).filter_by(request_id=new["id"])} >= {"shown", "used", "ignored"}


# ======================= US18 Retraining & model management =======================
def _register_active_v12(db, macro_f1: float) -> ModelRegistry:
    clf, metrics = classifier.load_or_train()
    path = training.artifact_dir() / "classifier" / "v1.2" / "model.joblib"
    path.parent.mkdir(parents=True, exist_ok=True)
    joblib.dump({"pipeline": clf.pipeline, "version": "v1.2", "metrics": metrics}, path)
    reg = ModelRegistry(version="v1.2", model_type="classifier", algorithm=classifier.ALGORITHM, accuracy=0.86,
                        macro_f1=macro_f1, metrics={}, status="active", is_active=True,
                        artifact_path="classifier/v1.2/model.joblib", trained_at=utcnow() - timedelta(days=7))
    db.add(reg)
    db.commit()
    training.sync_active_model(db)
    return reg


def _add_feedback(client, db, n: int) -> None:
    req, _ = _submit(client, "the window in my room does not close")
    agent = db.query(User).filter_by(email="agent1.housing@qoldau.edu.kz").one()
    samples = dataset.generate(seed=7)[:n]
    for text, label in samples:
        db.add(AIFeedback(request_id=req["id"], text=text, predicted_category="Other", corrected_category=label,
                          agent_id=agent.id, reason="wrong category"))
    db.commit()


def test_us18_retrain_from_feedback_auto_promotes(client, db, restore_model):
    """Given 250 new ai_feedback records and active v1.2 with macro-F1 0.85, When the admin clicks Retrain,
    Then a job trains, evaluates and registers v1.3 with metrics + confusion matrix, And v1.3 is activated
    because its macro-F1 >= 0.85, and the switch is logged."""
    v12 = _register_active_v12(db, 0.85)
    _add_feedback(client, db, 250)
    h = login(client, "admin@qoldau.edu.kz")
    assert client.get(f"{API}/ai/models/overview", headers=h).json()["feedback_since_last_training"] == 250
    r = client.post(f"{API}/ai/retrain", headers=h, json={})
    assert r.status_code == 202, r.text
    run = client.get(f"{API}/ai/training-runs/{r.json()['id']}", headers=h).json()
    assert run["status"] == "done", run["logs"]
    assert run["result"]["classifier"]["version"] == "v1.3" and run["result"]["classifier"]["promoted"]

    v13 = db.query(ModelRegistry).filter_by(model_type="classifier", version="v1.3").one()
    detail = client.get(f"{API}/ai/models/{v13.id}", headers=h).json()
    assert detail["is_active"] and detail["macro_f1"] >= 0.85 and detail["feedback_count"] == 250
    assert len(detail["metrics"]["confusion_matrix"]["matrix"]) == len(detail["metrics"]["confusion_matrix"]["labels"])
    assert set(detail["metrics"]["per_class_f1"]) == set(dataset.CATEGORIES)
    db.refresh(v12)
    assert v12.status == "archived" and not v12.is_active
    assert db.query(AuditLog).filter_by(entity="model", action="promote", entity_id=str(v13.id)).count() == 1
    new, _ = _submit(client, "I need a transcript of records for my visa application")
    assert new["id"] and db.get(Request, new["id"]).model_version == "v1.3"


def test_us18_lower_f1_model_is_not_activated(client, db, restore_model):
    _register_active_v12(db, 0.9999)
    _add_feedback(client, db, 30)
    h = login(client, "admin@qoldau.edu.kz")
    r = client.post(f"{API}/ai/retrain", headers=h, json={})
    assert r.status_code == 409 and "30/200" in r.json()["detail"]["message"]
    r = client.post(f"{API}/ai/retrain", headers=h, json={"force": True})
    assert r.status_code == 202
    v13 = db.query(ModelRegistry).filter_by(model_type="classifier", version="v1.3").one()
    assert v13.status == "candidate" and not v13.is_active
    assert db.query(ModelRegistry).filter_by(model_type="classifier", is_active=True).one().version == "v1.2"
    assert classifier.current_version() == "v1.2"


def test_us18_rollback_without_downtime(client, db, restore_model):
    """Given v1.3 is active, When the admin selects v1.2 and clicks Rollback, Then v1.2 becomes active within
    1 min without downtime, And new requests store model_version = v1.2."""
    v12 = _register_active_v12(db, 0.85)
    _add_feedback(client, db, 210)
    h = login(client, "admin@qoldau.edu.kz")
    client.post(f"{API}/ai/retrain", headers=h, json={})
    assert classifier.current_version() == "v1.3"
    r = client.post(f"{API}/ai/models/{v12.id}/rollback", headers=h)
    assert r.status_code == 200 and r.json()["is_active"]
    assert classifier.current_version() == "v1.2"  # swapped in this worker immediately
    assert db.query(AuditLog).filter_by(entity="model", action="rollback").count() == 1
    req, _ = _submit(client, "The radiator in room 412 of dorm 3 is leaking")
    assert db.get(Request, req["id"]).model_version == "v1.2"
    # Another worker still serving v1.3 picks the rollback up on its next registry sync.
    classifier.swap(*classifier.load_artifact(training.artifact_dir() / "classifier" / "v1.3" / "model.joblib"))
    assert training.sync_active_model(db) == "v1.2"
    assert client.post(f"{API}/ai/retrain", headers=login(client, "head.it@qoldau.edu.kz"), json={}).status_code == 403


# ======================= US17 Report export =======================
def _it_and_finance_requests(client, db):
    set_online(db, "agent1.it@qoldau.edu.kz", "agent1.finance@qoldau.edu.kz")
    for text in ("I cannot log in to Moodle, password reset does not work", "my university email account is locked",
                 "Wi-Fi does not work in the library", "Scholarship was not paid this month",
                 "I want to pay tuition in instalments"):
        _submit(client, text)


def test_us17_admin_exports_analytics_to_excel(client, db):
    """Given the admin filtered the dashboard to a period and Unit = IT Help Desk, When they click Export -> Excel,
    Then within 30 s an .xlsx with the 6 sheets is ready, And totals equal the dashboard and the export is audited."""
    _it_and_finance_requests(client, db)
    it = db.query(Unit).filter_by(name="IT Help Desk").one()
    today = utcnow().date().isoformat()
    h = login(client, "admin@qoldau.edu.kz")
    dash = client.get(f"{API}/analytics/dashboard?date_from={today}&date_to={today}&unit_id={it.id}", headers=h).json()
    t = time.perf_counter()
    r = client.post(f"{API}/export/excel", headers=h, json={"date_from": today, "date_to": today, "unit_id": it.id})
    assert r.status_code == 202, r.text
    job = client.get(f"{API}/export/jobs/{r.json()['id']}", headers=h).json()
    assert job["status"] == "done" and time.perf_counter() - t < 30
    f = client.get(job["download_url"])
    assert f.status_code == 200 and "spreadsheetml" in f.headers["content-type"]
    wb = load_workbook(io.BytesIO(f.content))
    assert wb.sheetnames == ["Summary", "By Category", "By Unit", "SLA", "AI Accuracy", "Raw Requests"]
    summary = {row[0]: row[1] for row in wb["Summary"].iter_rows(values_only=True) if row[0]}
    assert summary["Total requests"] == dash["kpis"]["total"] == 3
    assert summary["Open"] == dash["kpis"]["open"]
    assert wb["Raw Requests"].max_row - 1 == dash["kpis"]["total"]
    assert {row[4] for row in wb["Raw Requests"].iter_rows(min_row=2, values_only=True)} == {"IT Help Desk"}
    emails = [row[10] for row in wb["Raw Requests"].iter_rows(min_row=2, values_only=True)]
    assert all("@" in e and "***" not in e for e in emails)  # admins see contact data
    assert db.query(AuditLog).filter_by(entity="export", entity_id=str(job["id"])).count() == 1
    assert client.get(job["download_url"].replace("sig=", "sig=0")).status_code == 403


def test_us17_unit_head_pdf_is_scoped_and_students_get_403(client, db):
    """Given the Finance unit head, When they export a PDF for the current month, Then it contains only Finance
    Office requests, And a student calling POST /export/pdf receives HTTP 403."""
    _it_and_finance_requests(client, db)
    fin = db.query(Unit).filter_by(name="Finance Office").one()
    it = db.query(Unit).filter_by(name="IT Help Desk").one()
    h = login(client, "head.finance@qoldau.edu.kz")
    first = utcnow().date().replace(day=1).isoformat()
    r = client.post(f"{API}/export/pdf", headers=h, json={"date_from": first, "unit_id": it.id})  # tries another unit
    assert r.status_code == 202
    job = client.get(f"{API}/export/jobs/{r.json()['id']}", headers=h).json()
    assert job["status"] == "done" and job["filters"]["unit_id"] == fin.id
    finance_total = db.query(Request).filter_by(unit_id=fin.id).count()
    assert job["row_count"] == finance_total == 2
    pdf = client.get(job["download_url"])
    assert pdf.content.startswith(b"%PDF") and pdf.headers["content-type"] == "application/pdf"
    xl = client.post(f"{API}/export/excel", headers=h, json={"unit_id": it.id}).json()
    wb = load_workbook(io.BytesIO(client.get(client.get(f"{API}/export/jobs/{xl['id']}", headers=h).json()["download_url"]).content))
    rows = list(wb["Raw Requests"].iter_rows(min_row=2, values_only=True))
    assert {r[4] for r in rows} == {"Finance Office"} and all("***" in r[10] for r in rows)  # e-mails masked
    for email in ("student@qoldau.edu.kz", "agent1.it@qoldau.edu.kz"):
        assert client.post(f"{API}/export/pdf", headers=login(client, email), json={}).status_code == 403
    # Someone else's job is invisible.
    assert client.get(f"{API}/export/jobs/{job['id']}", headers=login(client, "head.it@qoldau.edu.kz")).status_code == 404


# ======================= embeddings & PII (prompt 4) =======================
def test_cross_language_queries_find_the_same_faq(client, kb):
    t = time.perf_counter()
    for query in ("how to order a transcript of records", "как заказать транскрипт", "транскриптке тапсырыс беру"):
        for lang in (embeddings.detect_language(query), None):
            hits = embeddings.search(kb, query, "faq", top_k=3, language=lang)
            slugs = [kb.get(FAQArticle, int(h.entity_id)).slug for h in hits]
            assert "ordering-transcript" in slugs, (query, lang, slugs)
    assert (time.perf_counter() - t) / 6 < 0.3


def test_language_detection():
    assert embeddings.detect_language("Не могу войти в Moodle") == "ru"
    assert embeddings.detect_language("Оқу орнынан анықтама керек") == "kz"
    assert embeddings.detect_language("I need a transcript") == "en"


def test_pii_masker():
    text = ("My name is Aruzhan Bekova, IIN 030512450123, phone +7 (701) 555-12-34, mail a.bekova@sdu.edu.kz, "
            "request REQ-2026-00123. Меня зовут Дильназ. Менің атым Әлихан.")
    out = pii.mask(text, ["Aruzhan Bekova"])
    for secret in ("Aruzhan", "Bekova", "030512450123", "701", "a.bekova@", "REQ-2026-00123", "Дильназ", "Әлихан"):
        assert secret not in out, secret
    assert "[email]" in out and "[phone]" in out and "[id]" in out and "[name]" in out
    assert pii.mask("Room 412 in dorm 3, deadline 25.10") == "Room 412 in dorm 3, deadline 25.10"
