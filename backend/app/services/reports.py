"""Report export to Excel and PDF (US17).

Every number comes from analytics.compute(), the function behind the dashboard, so a report always
matches what the user saw on screen. Unit heads are scoped to their own unit here as well as in the
API. Student e-mails are masked for everyone except Admins. Files are written privately to
settings.export_dir (stand-in for S3) and downloaded through signed links valid 24 h.
"""
from __future__ import annotations

import logging
import uuid
from datetime import date, datetime, timedelta
from io import BytesIO
from pathlib import Path

from sqlalchemy.orm import Session, joinedload

from ..config import get_settings
from ..db import SessionLocal, utcnow
from ..models import Category, ReportJob, Request, Role, Unit, User
from . import notifications

log = logging.getLogger(__name__)

BRAND = "1B45B4"  # SDU blue, the app's action colour
SHEETS = ("Summary", "By Category", "By Unit", "SLA", "AI Accuracy", "Raw Requests")
RAW_HEADERS = ["Request ID", "Created", "Subject", "Category", "Unit", "Priority", "Score", "Status", "Assignee",
               "Student", "Student email", "First response", "Resolved", "SLA due", "SLA breached", "CSAT",
               "Model version"]


def filters_from(job: ReportJob, requester: User):
    from ..api.analytics import Filters
    d = dict(job.filters or {})
    f = Filters(date_from=date.fromisoformat(d["date_from"]) if d.get("date_from") else None,
                date_to=date.fromisoformat(d["date_to"]) if d.get("date_to") else None,
                unit_id=d.get("unit_id"), category_id=d.get("category_id"), days=d.get("days"),
                priority=d.get("priority"), status=d.get("status"))
    if requester.role != Role.ADMIN:
        f.unit_id = requester.unit_id  # enforced in the query, not only in the UI
    if f.days and not f.date_from:
        f.date_from = (utcnow() - timedelta(days=f.days - 1)).date()
    return f


def mask_email(email: str) -> str:
    local, _, domain = email.partition("@")
    return f"{local[:1]}***@{domain}" if domain else "***"


def describe_filters(db: Session, f) -> list[tuple[str, str]]:
    unit = db.get(Unit, f.unit_id).name if f.unit_id and db.get(Unit, f.unit_id) else "All units"
    cat = db.get(Category, f.category_id).name if f.category_id and db.get(Category, f.category_id) else "All categories"
    period = f"{f.date_from or 'start'} - {f.date_to or 'today'}"
    return [("Period", period), ("Unit", unit), ("Category", cat), ("Priority", f.priority or "All"),
            ("Status", f.status or "All")]


def kpi_rows(k: dict) -> list[tuple[str, object]]:
    return [
        ("Total requests", k["total"]), ("Open", k["open"]), ("Resolved", k["resolved"]),
        ("Avg first response (h)", k["avg_first_response_hours"]), ("Avg resolution (h)", k["avg_resolution_hours"]),
        ("SLA compliance (%)", k["sla_compliance"]), ("CSAT average (1-5)", k["csat_avg"]),
        ("CSAT responses", k["csat_count"]), ("AI accuracy (%)", k["ai_accuracy"]), ("AI overrides", k["overrides"]),
        ("Self-service deflections", k.get("deflections")), ("Deflection rate (%)", k.get("deflection_rate")),
        ("Smart-reply usage (%)", k.get("smart_reply_usage")),
    ]


def raw_rows(db: Session, f, viewer: User) -> list[list]:
    from ..api.analytics import base_query
    admin = viewer.role == Role.ADMIN
    rows = (base_query(db, f).options(joinedload(Request.category), joinedload(Request.unit),
                                      joinedload(Request.assignee), joinedload(Request.student),
                                      joinedload(Request.csat))
            .order_by(Request.created_at).all())
    out = []
    for r in rows:
        email = r.student.email if r.student else ""
        out.append([r.id, r.created_at, r.subject, r.category.name if r.category else "", r.unit.name if r.unit else "",
                    r.priority, r.priority_score, r.status, r.assignee.name if r.assignee else "",
                    r.student.name if r.student else "", email if admin else mask_email(email),
                    r.first_response_at, r.resolved_at, r.sla_resolution_due, "yes" if r.sla_breached else "no",
                    r.csat.rating if r.csat else None, r.model_version or ""])
    return out


# ---------------- Excel ----------------
def build_excel(db: Session, f, viewer: User) -> tuple[bytes, int]:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    from ..api.analytics import compute
    data = compute(db, f)
    k = data["kpis"]
    wb = Workbook()
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=BRAND)

    def sheet(title: str, headers: list[str], rows: list[list], widths: dict[int, int] | None = None, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill, c.alignment = head_font, head_fill, Alignment(vertical="center")
        for row in rows:
            ws.append(row)
        ws.freeze_panes = "A2"
        for i, h in enumerate(headers, start=1):
            col = get_column_letter(i)
            sample = [len(str(r[i - 1])) for r in rows[:500] if r[i - 1] is not None]
            ws.column_dimensions[col].width = min(60, max([len(h), *sample, 8]) + 2)
            for cell in ws[col][1:]:
                if isinstance(cell.value, datetime):
                    cell.number_format = "yyyy-mm-dd hh:mm"
                elif isinstance(cell.value, float):
                    cell.number_format = "0.0"
        return ws

    summary = sheet("Summary", ["Metric", "Value"], [[n, v] for n, v in kpi_rows(k)], first=True)
    summary.insert_rows(1, amount=len(describe_filters(db, f)) + 3)
    summary["A1"], summary["A1"].font = "Qoldau - Analytics report (SDU University)", Font(bold=True, size=14, color=BRAND)
    summary["A2"] = f"Generated {utcnow():%Y-%m-%d %H:%M} UTC by {viewer.name}"
    for i, (name, val) in enumerate(describe_filters(db, f), start=3):
        summary[f"A{i}"], summary[f"B{i}"] = name, val
    summary.freeze_panes = None

    total = k["total"] or 1
    sheet("By Category", ["Category", "Requests", "Share (%)"],
          [[c["category"], c["count"], round(100 * c["count"] / total, 1)] for c in data["by_category"]])
    sheet("By Unit", ["Unit", "New", "In progress", "Waiting for student", "Resolved", "Total"],
          [[u["unit"], u["new"], u["in_progress"], u["waiting"], u["resolved"],
            u["new"] + u["in_progress"] + u["waiting"] + u["resolved"]] for u in data["by_unit"]])
    sheet("SLA", ["Priority", "Requests with decided SLA", "Compliance (%)"],
          [[s["priority"], s["decided"], s["compliance"]] for s in data["sla_by_priority"]]
          + [["All", sum(s["decided"] for s in data["sla_by_priority"]), k["sla_compliance"]]])
    sheet("AI Accuracy", ["Category", "Classified", "Overridden by agents", "Accuracy (%)"],
          [[a["category"], a["classified"], a["overridden"], a["accuracy"]] for a in data["ai_by_category"]]
          + [["All", sum(a["classified"] for a in data["ai_by_category"]), k["overrides"], k["ai_accuracy"]]])
    raw = raw_rows(db, f, viewer)
    sheet("Raw Requests", RAW_HEADERS, raw)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue(), len(raw)


# ---------------- PDF ----------------
_FONT: tuple[str, str] | None = None


def _fonts() -> tuple[str, str]:
    """A TTF with Cyrillic / Kazakh glyphs when one is available (Helvetica has none)."""
    global _FONT
    if _FONT:
        return _FONT
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    candidates = [("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
                  ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf")]
    try:
        import matplotlib
        base = Path(matplotlib.get_data_path()) / "fonts" / "ttf"
        candidates.insert(0, (str(base / "DejaVuSans.ttf"), str(base / "DejaVuSans-Bold.ttf")))
    except Exception:
        pass
    for regular, bold in candidates:
        if Path(regular).exists() and Path(bold).exists():
            pdfmetrics.registerFont(TTFont("QoldauSans", regular))
            pdfmetrics.registerFont(TTFont("QoldauSans-Bold", bold))
            _FONT = ("QoldauSans", "QoldauSans-Bold")
            return _FONT
    _FONT = ("Helvetica", "Helvetica-Bold")
    return _FONT


def build_pdf(db: Session, f, viewer: User) -> tuple[bytes, int]:
    from reportlab.graphics.charts.barcharts import HorizontalBarChart
    from reportlab.graphics.charts.lineplots import LinePlot
    from reportlab.graphics.shapes import Drawing, String
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    from ..api.analytics import compute
    regular, bold = _fonts()
    data = compute(db, f)
    k = data["kpis"]
    brand = colors.HexColor(f"#{BRAND}")
    series = [colors.HexColor(c) for c in ("#1B45B4", "#0F9D8A", "#E0A100", "#C2410C", "#7C3AED", "#64748B")]
    h1 = ParagraphStyle("h1", fontName=bold, fontSize=22, leading=28, textColor=brand)
    h2 = ParagraphStyle("h2", fontName=bold, fontSize=14, leading=20, spaceBefore=10, spaceAfter=6)
    body = ParagraphStyle("body", fontName=regular, fontSize=10, leading=14)
    muted = ParagraphStyle("muted", parent=body, textColor=colors.HexColor("#64748B"))

    def table(rows, widths, header=True):
        t = Table(rows, colWidths=widths, repeatRows=1 if header else 0)
        style = [("FONT", (0, 0), (-1, -1), regular, 9), ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
                 ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
                 ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]
        if header:
            style += [("BACKGROUND", (0, 0), (-1, 0), brand), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                      ("FONT", (0, 0), (-1, 0), bold, 9)]
        t.setStyle(TableStyle(style))
        return t

    fmt = lambda v: "-" if v is None else (f"{v:.1f}" if isinstance(v, float) else str(v))  # noqa: E731
    story = [Spacer(1, 30 * mm), Paragraph("Qoldau", h1),
             Paragraph("Student Request Routing &amp; Support - analytics report", h2),
             Paragraph("SDU University", muted), Spacer(1, 8 * mm),
             table([[n, v] for n, v in describe_filters(db, f)], [45 * mm, 110 * mm], header=False),
             Spacer(1, 4 * mm),
             Paragraph(f"Generated {utcnow():%Y-%m-%d %H:%M} UTC by {viewer.name}", muted), PageBreak(),
             Paragraph("Key indicators", h2),
             table([["Metric", "Value"]] + [[n, fmt(v)] for n, v in kpi_rows(k)], [100 * mm, 55 * mm])]

    # Requests over time (created vs resolved)
    pts = data["over_time"]
    if len(pts) > 1:
        d = Drawing(170 * mm, 70 * mm)
        lp = LinePlot()
        lp.x, lp.y, lp.width, lp.height = 12 * mm, 10 * mm, 150 * mm, 52 * mm
        lp.data = [[(i, p["created"]) for i, p in enumerate(pts)], [(i, p["resolved"]) for i, p in enumerate(pts)]]
        for i in range(2):
            lp.lines[i].strokeColor, lp.lines[i].strokeWidth = series[i], 1.5
        lp.xValueAxis.visibleLabels = False
        lp.yValueAxis.labels.fontName = regular
        d.add(lp)
        d.add(String(12 * mm, 65 * mm, f"Requests over time: created (blue) vs resolved (teal), {pts[0]['date']} - "
                                        f"{pts[-1]['date']}", fontName=regular, fontSize=9))
        story += [KeepTogether([Paragraph("Volume", h2), d])]

    cats = data["by_category"][:10]
    if cats:
        d = Drawing(170 * mm, max(40, 9 * len(cats)) * mm)
        bc = HorizontalBarChart()
        bc.x, bc.y, bc.width, bc.height = 55 * mm, 5 * mm, 105 * mm, (max(40, 9 * len(cats)) - 10) * mm
        bc.data = [[c["count"] for c in reversed(cats)]]
        bc.categoryAxis.categoryNames = [c["category"] for c in reversed(cats)]
        bc.categoryAxis.labels.fontName = bc.valueAxis.labels.fontName = regular
        bc.categoryAxis.labels.fontSize = 8
        bc.bars[0].fillColor = brand
        bc.valueAxis.valueMin = 0
        d.add(bc)
        story += [KeepTogether([Paragraph("Requests by category", h2), d])]

    prio = list(reversed(data["priority"]))  # Low ... Critical, drawn bottom-up
    if any(p["count"] for p in prio):
        # Same priority colours as the app's badges; each bar is labelled, so colour is never the only cue.
        prio_colors = {"Critical": "#C42020", "High": "#C2410C", "Medium": "#8A5A00", "Low": "#3D6B12"}
        d = Drawing(170 * mm, 42 * mm)
        bc = HorizontalBarChart()
        bc.x, bc.y, bc.width, bc.height = 30 * mm, 4 * mm, 130 * mm, 34 * mm
        bc.data = [[p["count"] for p in prio]]
        bc.categoryAxis.categoryNames = [f"{p['priority']} ({p['count']})" for p in prio]
        bc.categoryAxis.labels.fontName = bc.valueAxis.labels.fontName = regular
        bc.valueAxis.valueMin = 0
        for i, p in enumerate(prio):
            bc.bars[(0, i)].fillColor = colors.HexColor(prio_colors.get(p["priority"], "#64748B"))
        d.add(bc)
        story += [KeepTogether([Paragraph("Priority distribution", h2), d])]

    # Each heading stays on the same page as its table.
    story += [KeepTogether([
                  Paragraph("Workload by unit", h2),
                  table([["Unit", "New", "In progress", "Waiting", "Resolved"]]
                        + [[u["unit"], u["new"], u["in_progress"], u["waiting"], u["resolved"]] for u in data["by_unit"]],
                        [70 * mm, 20 * mm, 25 * mm, 20 * mm, 20 * mm])]),
              KeepTogether([
                  Paragraph("SLA compliance", h2),
                  table([["Priority", "Decided", "Compliance (%)"]]
                        + [[s["priority"], s["decided"], fmt(s["compliance"])] for s in data["sla_by_priority"]]
                        + [["All", sum(s["decided"] for s in data["sla_by_priority"]), fmt(k["sla_compliance"])]],
                        [60 * mm, 40 * mm, 55 * mm])]),
              KeepTogether([
                  Paragraph("AI accuracy", h2),
                  Paragraph("Share of classified requests whose category agents did not have to correct.", muted),
                  Spacer(1, 2 * mm),
                  table([["Category", "Classified", "Overridden", "Accuracy (%)"]]
                        + [[a["category"], a["classified"], a["overridden"], fmt(a["accuracy"])] for a in data["ai_by_category"]]
                        + [["All", sum(a["classified"] for a in data["ai_by_category"]), k["overrides"], fmt(k["ai_accuracy"])]],
                        [65 * mm, 30 * mm, 30 * mm, 30 * mm])])]

    def chrome(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(brand)
        canvas.rect(0, A4[1] - 8 * mm, A4[0], 8 * mm, stroke=0, fill=1)
        canvas.setFont(regular, 8)
        canvas.setFillColor(colors.HexColor("#64748B"))
        canvas.drawString(18 * mm, 10 * mm, "Qoldau - SDU University - confidential")
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, f"Page {doc.page}")
        canvas.restoreState()

    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=18 * mm,
                            bottomMargin=18 * mm, title="Qoldau analytics report", author=viewer.name)
    doc.build(story, onFirstPage=chrome, onLaterPages=chrome)
    return buf.getvalue(), k["total"]


# ---------------- job ----------------
def run_export(job_id: int) -> None:
    db = SessionLocal()
    try:
        job = db.get(ReportJob, job_id)
        if job is None:
            return
        user = db.get(User, job.requested_by)
        job.status = "running"
        db.commit()
        try:
            f = filters_from(job, user)
            content, rows = (build_excel if job.type == "excel" else build_pdf)(db, f, user)
            ext = "xlsx" if job.type == "excel" else "pdf"
            key = f"{user.id}/{utcnow():%Y%m%d-%H%M%S}-{uuid.uuid4().hex[:8]}.{ext}"
            path = get_settings().export_dir / key
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(content)
            job.status, job.file_key, job.file_size, job.row_count = "done", key, len(content), rows
            job.finished_at = utcnow()
            notifications.notify(db, user, "report_ready", f"Your {job.type.upper()} report is ready",
                                 f"{rows} requests. Download it from Reports within {get_settings().export_link_hours} h.",
                                 email=False)
        except Exception as exc:
            log.exception("export %s failed", job_id)
            db.rollback()
            job = db.get(ReportJob, job_id)
            job.status, job.error, job.finished_at = "failed", f"{type(exc).__name__}: {exc}", utcnow()
            notifications.notify(db, user, "report_failed", f"{job.type.upper()} report failed", job.error, email=False)
        db.commit()
    finally:
        db.close()


def file_path(job: ReportJob) -> Path:
    return get_settings().export_dir / job.file_key
